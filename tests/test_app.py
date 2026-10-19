"""Тесты Flask-приложения: /, /health, /upload, /api/anonymize."""
import io
import json

import pytest

from app import app as flask_app


@pytest.fixture
def client():
    flask_app.config["TESTING"] = True
    with flask_app.test_client() as c:
        yield c


def test_index_serves_ui(client):
    r = client.get("/")
    assert r.status_code == 200
    assert "Анонимайзер" in r.get_data(as_text=True)


def test_health(client):
    r = client.get("/health")
    assert r.status_code == 200
    assert r.get_json()["status"] == "healthy"


def _upload(client, filename, content, **form):
    data = {"file": (io.BytesIO(content), filename)}
    data.update(form)
    return client.post("/upload", data=data, content_type="multipart/form-data")


def test_upload_anonymizes_file(client):
    src = "user Иванов Иван Иванович from 10.0.0.1:8080 tel +79991234567\n".encode()
    r = _upload(client, "test.log", src)
    assert r.status_code == 200
    text = r.get_data(as_text=True)
    assert "10.0.0.1" not in text and "8080" not in text
    assert "Иванов" not in text and "+79991234567" not in text
    assert "[IP1]" in text and "[PORT1]" in text
    assert "test_clean" in r.headers["Content-Disposition"]
    stats = json.loads(r.headers["X-Anon-Stats"])
    assert stats["total_replacements"] > 0


def test_upload_no_file(client):
    r = client.post("/upload", data={}, content_type="multipart/form-data")
    assert r.status_code == 400


def test_upload_wrong_extension_rejected(client):
    r = _upload(client, "evil.exe", b"MZ\x00\x00binary")
    assert r.status_code == 400


def test_upload_binary_rejected(client):
    r = _upload(client, "data.log", bytes(range(256)) * 40)
    assert r.status_code == 400


def test_upload_cp1251_encoding(client):
    src = "ФИО: Петров Пётр Петрович\n".encode("cp1251")
    r = _upload(client, "win.log", src)
    assert r.status_code == 200
    assert "Петров" not in r.get_data(as_text=True)


def test_upload_entities_filter(client):
    src = "ip 10.0.0.1 tel +79991234567 host example.com\n".encode()
    r = _upload(client, "f.log", src, entities="phone")
    text = r.get_data(as_text=True)
    assert "10.0.0.1" in text          # ip отключен
    assert "+79991234567" not in text  # phone работает


def test_upload_entities_multiple_values(client):
    """UI шлёт каждое поле entities отдельным значением - все должны учитываться."""
    src = "ip 10.0.0.1 tel +79991234567 host example.com\n".encode()
    r = _upload(client, "f.log", src, entities=["phone", "host"])
    text = r.get_data(as_text=True)
    assert "10.0.0.1" in text           # ip не выбран
    assert "+79991234567" not in text   # phone выбран
    assert "example.com" not in text    # host выбран


def test_api_anonymize_entities_list(client):
    r = client.post("/api/anonymize", json={
        "text": "10.0.0.1 +79991234567", "entities": ["phone"],
    })
    assert r.status_code == 200
    assert "10.0.0.1" in r.get_json()["text"]
    assert "+79991234567" not in r.get_json()["text"]


def test_upload_style_and_mode(client):
    src = b"a 10.0.0.1 b 10.0.0.1\n"
    r1 = _upload(client, "a.log", src, style="realistic")
    assert "10.0.0.1" not in r1.get_data(as_text=True)
    # stable-режим с солью воспроизводим между запросами
    r2 = _upload(client, "b.log", src, mode="stable", salt="s1")
    r3 = _upload(client, "c.log", src, mode="stable", salt="s1")
    assert r2.get_data() == r3.get_data()
    r4 = _upload(client, "d.log", src, mode="stable", salt="s2")
    assert r2.get_data() != r4.get_data()


def test_upload_bad_entities_400(client):
    r = _upload(client, "f.log", b"x 1.2.3.4\n", entities="ip,unknown")
    assert r.status_code == 400


def test_api_anonymize_json(client):
    r = client.post("/api/anonymize", json={
        "text": "Иванов Иван Иванович 10.0.0.1:22 mail.ru",
    })
    assert r.status_code == 200
    data = r.get_json()
    assert "10.0.0.1" not in data["text"]
    assert "Иванов" not in data["text"]
    assert data["mapping"]
    assert data["stats"]["changed_lines"] == 1


def test_api_anonymize_requires_text(client):
    assert client.post("/api/anonymize", json={}).status_code == 400
    assert client.post("/api/anonymize", data="junk").status_code == 400


# ------------------------------------------------------------ Локализация


def test_upload_error_localized_en(client):
    r = client.post("/upload", data={"lang": "en"},
                    content_type="multipart/form-data")
    assert r.status_code == 400
    assert r.get_json()["error"] == "No file provided"


def test_upload_error_localized_ru(client):
    r = client.post("/upload", data={"lang": "ru"},
                    content_type="multipart/form-data")
    assert r.status_code == 400
    assert r.get_json()["error"] == "Файл не передан"


def test_upload_not_allowed_localized(client):
    r = _upload(client, "evil.exe", b"MZ\x00\x00binary", lang="en")
    assert r.status_code == 400
    assert "Allowed" in r.get_json()["error"]
    r2 = _upload(client, "evil.exe", b"MZ\x00\x00binary", lang="ru")
    assert "Разрешены" in r2.get_json()["error"]


def test_api_error_localized_en(client):
    r = client.post("/api/anonymize", json={"lang": "en"})
    assert r.status_code == 400
    assert "text" in r.get_json()["error"]


def test_api_error_localized_ru(client):
    r = client.post("/api/anonymize", json={"lang": "ru"})
    assert r.status_code == 400
    assert "text" in r.get_json()["error"] and "обязательно" in r.get_json()["error"]


def test_corrupt_xlsx_error_localized(client):
    r = _upload(client, "bad.xlsx", b"PK\x03\x04 garbage", lang="ru")
    assert r.status_code == 400
    assert "Excel" in r.get_json()["error"]
    r2 = _upload(client, "bad.xlsx", b"PK\x03\x04 garbage", lang="en")
    assert "Excel" in r2.get_json()["error"]


# ---------------------------------------------------------------- Excel


def _make_xlsx(rows) -> bytes:
    import io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_xls(rows) -> bytes:
    import io
    import xlwt
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("S")
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            ws.write(r, c, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_upload_xlsx(client):
    import io
    import openpyxl
    raw = _make_xlsx([
        ["Клиент", "IP"],
        ["Иванов Иван Иванович", "10.0.0.1:8080"],
    ])
    r = _upload(client, "report.xlsx", raw)
    assert r.status_code == 200
    assert "report_clean" in r.headers["Content-Disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    vals = [[c.value for c in row] for row in wb.active.iter_rows()]
    assert "Иванов" not in str(vals)
    assert "10.0.0.1" not in str(vals)
    assert "[IP1]" in str(vals)


def test_upload_xls(client):
    import xlrd
    raw = _make_xls([["tel", "+79991234567"]])
    r = _upload(client, "old.xls", raw)
    assert r.status_code == 200
    book = xlrd.open_workbook(file_contents=r.data)
    vals = str(book.sheet_by_index(0).row_values(0))
    assert "+79991234567" not in vals
    assert "[PHONE1]" in vals


def test_upload_corrupt_xlsx_rejected(client):
    r = _upload(client, "bad.xlsx", b"PK\x03\x04 garbage not a real xlsx")
    assert r.status_code == 400


def test_upload_xlsx_entities_filter(client):
    raw = _make_xlsx([["10.0.0.1", "+79991234567"]])
    r = _upload(client, "f.xlsx", raw, entities="phone")
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    vals = [[c.value for c in row] for row in wb.active.iter_rows()][0]
    assert vals[0] == "10.0.0.1"          # ip не выбран
    assert "+79991234567" not in str(vals[1])
