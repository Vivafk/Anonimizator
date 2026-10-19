"""Тесты анонимизации Excel-файлов (.xls/.xlsx)."""
import io

import pytest
import openpyxl
import xlrd
import xlwt

from anonymizer import Anonymizer
from excel_anon import anonymize_excel


def make_xlsx(sheets: dict) -> bytes:
    """sheets: {имя_листа: [[значения строк], ...]}"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_xls(sheets: dict) -> bytes:
    wb = xlwt.Workbook(encoding="utf-8")
    for name, rows in sheets.items():
        ws = wb.add_sheet(name)
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                ws.write(r, c, val)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_xlsx(raw: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    return {ws.title: [[c.value for c in row] for row in ws.iter_rows()]
            for ws in wb.worksheets}


def read_xls(raw: bytes) -> dict:
    book = xlrd.open_workbook(file_contents=raw)
    return {sh.name: [sh.row_values(r) for r in range(sh.nrows)]
            for sh in book.sheets()}


# ------------------------------------------------------------------- xlsx

def test_xlsx_strings_anonymized():
    raw = make_xlsx({"Лист1": [
        ["Клиент", "Телефон", "IP"],
        ["Иванов Иван Иванович", "+79991234567", "10.0.0.1:8080"],
        ["Иванов Иван Иванович", "+79991234567", "10.0.0.2:22"],
    ]})
    out = anonymize_excel(raw, "xlsx", Anonymizer())[0]
    data = read_xlsx(out)
    rows = data["Лист1"]
    # заголовки не содержат чувствительных данных - не тронуты
    assert rows[0] == ["Клиент", "Телефон", "IP"]
    # чувствительные данные заменены
    assert "Иванов" not in str(rows[1])
    assert "+79991234567" not in str(rows[1])
    assert "10.0.0.1" not in str(rows[1])
    # одинаковые значения в разных строках -> одинаковые псевдонимы
    assert rows[1][0] == rows[2][0]
    assert rows[1][1] == rows[2][1]
    # разные IP -> разные псевдонимы
    assert rows[1][2] != rows[2][2]


def test_xlsx_multiple_sheets_share_mapping():
    raw = make_xlsx({
        "A": [["10.0.0.1 up"]],
        "B": [["10.0.0.1 again"]],
    })
    out = anonymize_excel(raw, "xlsx", Anonymizer())[0]
    data = read_xlsx(out)
    assert "[IP1]" in data["A"][0][0]
    assert "[IP1]" in data["B"][0][0]


def test_xlsx_numbers_and_dates_preserved():
    import datetime
    raw = make_xlsx({"S": [[42, 3.14, "text 10.0.0.1"]]})
    out = anonymize_excel(raw, "xlsx", Anonymizer())[0]
    row = read_xlsx(out)["S"][0]
    assert row[0] == 42
    assert row[1] == 3.14
    assert "10.0.0.1" not in str(row[2])


def test_xlsx_numeric_phone_replaced():
    """Телефон, сохранённый числом (89991234567), не должен протечь."""
    raw = make_xlsx({"S": [["tel", 89991234567], ["qty", 12345]]})
    out = anonymize_excel(raw, "xlsx", Anonymizer())[0]
    rows = read_xlsx(out)["S"]
    assert "9991234567" not in str(rows[0])
    assert "[PHONE1]" in str(rows[0])
    # обычное короткое число не тронуто
    assert rows[1][1] == 12345


def test_xlsx_sensitive_sheet_name_replaced():
    raw = make_xlsx({"Отчет Иванов": [["данные 10.0.0.1"]]})
    out = anonymize_excel(raw, "xlsx", Anonymizer())[0]
    data = read_xlsx(out)
    assert "Иванов" not in list(data.keys())[0]


def test_xlsx_empty_cells_and_none():
    raw = make_xlsx({"S": [[None, "10.0.0.1", None]]})
    out = anonymize_excel(raw, "xlsx", Anonymizer())[0]
    row = read_xlsx(out)["S"][0]
    assert row[0] is None
    assert row[1] == "[IP1]"


# -------------------------------------------------------------------- xls

def test_xls_strings_anonymized():
    raw = make_xls({"Sheet1": [
        ["ФИО", "Контакт"],
        ["Петров Пётр Петрович", "8-900-111-22-33"],
    ]})
    out = anonymize_excel(raw, "xls", Anonymizer())[0]
    rows = read_xls(out)["Sheet1"]
    assert "Петров" not in str(rows[1])
    assert "900-111-22-33" not in str(rows[1])
    assert "[PHONE1]" in str(rows[1])


def test_xls_numbers_preserved():
    raw = make_xls({"S": [[42, 3.14, "ip 10.0.0.5"]]})
    out = anonymize_excel(raw, "xls", Anonymizer())[0]
    row = read_xls(out)["S"][0]
    assert row[0] == 42.0
    assert row[1] == 3.14
    assert "10.0.0.5" not in str(row[2])


def test_xls_numeric_phone_replaced():
    raw = make_xls({"S": [[89991234567]]})
    out = anonymize_excel(raw, "xls", Anonymizer())[0]
    assert "[PHONE1]" in str(read_xls(out)["S"][0][0])


# ------------------------------------------------------------- ошибки

def test_invalid_xlsx_raises():
    with pytest.raises(ValueError):
        anonymize_excel(b"not an excel file at all", "xlsx", Anonymizer())


def test_invalid_xls_raises():
    with pytest.raises(ValueError):
        anonymize_excel(b"MZ\x00\x00junk", "xls", Anonymizer())


def test_unknown_extension_raises():
    raw = make_xlsx({"S": [["x"]]})
    with pytest.raises(ValueError):
        anonymize_excel(raw, "xlsm", Anonymizer())
